"""
Big-4 Style Report — Excel generation using openpyxl.
Produces a workbook with cover sheet, financial statements (with formulas),
strategic analysis, and audit trail.
"""
from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter


def _get_cover(payload: dict[str, Any]) -> dict[str, Any]:
    c = payload.get("cover") or {}
    if isinstance(c, dict):
        return c
    return {
        "title": getattr(c, "title", "Financial Report"),
        "subtitle": getattr(c, "subtitle", None),
        "entity_name": getattr(c, "entity_name", None),
        "report_date": getattr(c, "report_date", None),
        "period_label": getattr(c, "period_label", None),
        "prepared_by": getattr(c, "prepared_by", None),
        "codification": getattr(c, "codification", None),
    }


def _get_financials(payload: dict[str, Any]) -> dict[str, Any]:
    fs = payload.get("financial_statements") or {}
    if isinstance(fs, dict):
        return fs
    return {
        "balance_sheet": getattr(fs, "balance_sheet", {}),
        "profit_and_loss": getattr(fs, "profit_and_loss", {}),
        "report_date": getattr(fs, "report_date", None),
    }


def _fmt_num(value: Any) -> str | int | float:
    if value is None:
        return "—"
    try:
        n = float(value)
        return round(n, 2)
    except (TypeError, ValueError):
        return value


# --- Styles ---
HEADER_FILL = PatternFill(start_color="1a365d", end_color="1a365d", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=12)
SUBHEADER_FILL = PatternFill(start_color="e2e8f0", end_color="e2e8f0", fill_type="solid")
SUBHEADER_FONT = Font(bold=True, size=11)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def _write_cover_sheet(ws: Any, cover: dict[str, Any]) -> None:
    ws.title = "Cover"
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 40
    row = 1
    title = cover.get("title") or "Financial Report"
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    ws.cell(row=row, column=1, value=title)
    ws.cell(row=row, column=1).font = Font(bold=True, size=18, color="1a365d")
    ws.cell(row=row, column=1).alignment = Alignment(horizontal="center")
    row += 2
    if cover.get("subtitle"):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        ws.cell(row=row, column=1, value=cover["subtitle"])
        ws.cell(row=row, column=1).alignment = Alignment(horizontal="center")
        row += 1
    row += 1
    for label, key in [
        ("Entity", "entity_name"),
        ("Report date", "report_date"),
        ("Period", "period_label"),
        ("Prepared by", "prepared_by"),
        ("Codification", "codification"),
    ]:
        val = cover.get(key)
        if val is not None and val != "":
            ws.cell(row=row, column=1, value=label)
            ws.cell(row=row, column=1).font = Font(bold=True)
            ws.cell(row=row, column=2, value=val)
            row += 1
    row += 1
    ws.cell(row=row, column=1, value="CPA-verified financial statements · CFA-generated strategic analysis · Audit trail included")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    ws.cell(row=row, column=1).font = Font(italic=True, size=9, color="718096")


def _write_balance_sheet(ws: Any, bs: dict[str, Any], start_row: int) -> int:
    ws.cell(row=start_row, column=1, value="Balance Sheet")
    ws.cell(row=start_row, column=1).font = Font(bold=True, size=12)
    start_row += 1
    col_label, col_amount = 1, 2
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 14
    ws.cell(row=start_row, column=col_label, value="Account")
    ws.cell(row=start_row, column=col_amount, value="Amount")
    for c in (col_label, col_amount):
        ws.cell(row=start_row, column=c).fill = SUBHEADER_FILL
        ws.cell(row=start_row, column=c).font = SUBHEADER_FONT
        ws.cell(row=start_row, column=c).border = THIN_BORDER
    ws.cell(row=start_row, column=col_amount).alignment = Alignment(horizontal="right")
    start_row += 1
    first_data_row = start_row
    for key in ("assets", "liabilities", "equity"):
        arr = bs.get(key) or []
        for item in arr:
            label = item.get("label", "") if isinstance(item, dict) else str(item)
            amt = item.get("amount", 0) if isinstance(item, dict) else 0
            ws.cell(row=start_row, column=col_label, value=label)
            ws.cell(row=start_row, column=col_amount, value=float(amt) if amt is not None else None)
            ws.cell(row=start_row, column=col_amount).number_format = "#,##0.00"
            start_row += 1
        total_key = {"assets": "total_assets", "liabilities": "total_liabilities", "equity": "total_equity"}.get(key)
        total = bs.get(total_key) if total_key else None
        if total is not None:
            # Formula: sum of the block above for this section
            r_start = first_data_row
            r_end = start_row - 1
            if r_end >= r_start:
                ws.cell(row=start_row, column=col_label, value=f"Total {key.title()}")
                ws.cell(row=start_row, column=col_label).font = Font(bold=True)
                ws.cell(row=start_row, column=col_amount, value=f"=SUM(B{r_start}:B{r_end})")
                ws.cell(row=start_row, column=col_amount).font = Font(bold=True)
                ws.cell(row=start_row, column=col_amount).number_format = "#,##0.00"
                start_row += 1
        first_data_row = start_row
    return start_row


def _write_pl(ws: Any, pl: dict[str, Any], start_row: int) -> int:
    ws.cell(row=start_row, column=1, value="Statement of Comprehensive Income (P&L)")
    ws.cell(row=start_row, column=1).font = Font(bold=True, size=12)
    start_row += 1
    col_label, col_amount = 1, 2
    ws.cell(row=start_row, column=col_label, value="Account")
    ws.cell(row=start_row, column=col_amount, value="Amount")
    for c in (col_label, col_amount):
        ws.cell(row=start_row, column=c).fill = SUBHEADER_FILL
        ws.cell(row=start_row, column=c).font = SUBHEADER_FONT
        ws.cell(row=start_row, column=c).border = THIN_BORDER
    ws.cell(row=start_row, column=col_amount).alignment = Alignment(horizontal="right")
    start_row += 1
    rev_start = start_row
    for item in pl.get("revenue") or []:
        label = item.get("label", "") if isinstance(item, dict) else str(item)
        amt = item.get("amount", 0) if isinstance(item, dict) else 0
        ws.cell(row=start_row, column=col_label, value=label)
        ws.cell(row=start_row, column=col_amount, value=float(amt) if amt is not None else None)
        ws.cell(row=start_row, column=col_amount).number_format = "#,##0.00"
        start_row += 1
    rev_end = start_row - 1
    if rev_end >= rev_start and pl.get("total_revenue") is not None:
        ws.cell(row=start_row, column=col_label, value="Total Revenue")
        ws.cell(row=start_row, column=col_label).font = Font(bold=True)
        ws.cell(row=start_row, column=col_amount, value=f"=SUM(B{rev_start}:B{rev_end})")
        ws.cell(row=start_row, column=col_amount).font = Font(bold=True)
        ws.cell(row=start_row, column=col_amount).number_format = "#,##0.00"
        start_row += 1
    exp_start = start_row
    for item in pl.get("expenses") or []:
        label = item.get("label", "") if isinstance(item, dict) else str(item)
        amt = item.get("amount", 0) if isinstance(item, dict) else 0
        ws.cell(row=start_row, column=col_label, value=label)
        ws.cell(row=start_row, column=col_amount, value=float(amt) if amt is not None else None)
        ws.cell(row=start_row, column=col_amount).number_format = "#,##0.00"
        start_row += 1
    exp_end = start_row - 1
    if exp_end >= exp_start and pl.get("total_expenses") is not None:
        ws.cell(row=start_row, column=col_label, value="Total Expenses")
        ws.cell(row=start_row, column=col_label).font = Font(bold=True)
        ws.cell(row=start_row, column=col_amount, value=f"=SUM(B{exp_start}:B{exp_end})")
        ws.cell(row=start_row, column=col_amount).font = Font(bold=True)
        ws.cell(row=start_row, column=col_amount).number_format = "#,##0.00"
        start_row += 1
    if pl.get("net_income") is not None and rev_end >= rev_start and exp_end >= exp_start:
        total_rev_row = rev_end + 1
        total_exp_row = exp_end + 1
        ws.cell(row=start_row, column=col_label, value="Net Income")
        ws.cell(row=start_row, column=col_label).font = Font(bold=True)
        ws.cell(row=start_row, column=col_amount, value=f"=B{total_rev_row}-B{total_exp_row}")
        ws.cell(row=start_row, column=col_amount).font = Font(bold=True)
        ws.cell(row=start_row, column=col_amount).number_format = "#,##0.00"
        start_row += 1
    return start_row


def _write_financial_sheet(ws: Any, financials: dict[str, Any]) -> None:
    ws.title = "Financial Statements"
    row = 1
    ws.cell(row=row, column=1, value="CPA-Verified Financial Statements")
    ws.cell(row=row, column=1).font = Font(bold=True, size=14, color="1a365d")
    row += 1
    if financials.get("report_date"):
        ws.cell(row=row, column=1, value=f"As of {financials['report_date']}")
        ws.cell(row=row, column=1).font = Font(size=10, color="4a5568")
        row += 2
    bs = financials.get("balance_sheet") or {}
    if bs:
        row = _write_balance_sheet(ws, bs, row)
        row += 2
    pl = financials.get("profit_and_loss") or {}
    if pl:
        row = _write_pl(ws, pl, row)


def _write_strategic_sheet(ws: Any, strategic_analysis: Any) -> None:
    ws.title = "Strategic Analysis"
    ws.column_dimensions["A"].width = 100
    row = 1
    ws.cell(row=row, column=1, value="CFA-Generated Strategic Analysis")
    ws.cell(row=row, column=1).font = Font(bold=True, size=14, color="1a365d")
    row += 2
    if isinstance(strategic_analysis, str):
        for line in strategic_analysis.split("\n"):
            ws.cell(row=row, column=1, value=line.strip() or "")
            ws.cell(row=row, column=1).alignment = Alignment(wrap_text=True)
            row += 1
    elif isinstance(strategic_analysis, list):
        for item in strategic_analysis:
            if isinstance(item, dict):
                title = item.get("title", "")
                content = item.get("content", "")
                if title:
                    ws.cell(row=row, column=1, value=title)
                    ws.cell(row=row, column=1).font = Font(bold=True, size=11)
                    row += 1
                if content:
                    ws.cell(row=row, column=1, value=content)
                    ws.cell(row=row, column=1).alignment = Alignment(wrap_text=True)
                    row += 1


def _write_audit_sheet(ws: Any, audit_trail: list[Any]) -> None:
    ws.title = "Audit Trail"
    cols = ["Timestamp (UTC)", "Event type", "Reasoning", "Citations", "Outcome", "Warning"]
    for c, name in enumerate(cols, 1):
        ws.cell(row=1, column=c, value=name)
        ws.cell(row=1, column=c).fill = HEADER_FILL
        ws.cell(row=1, column=c).font = HEADER_FONT
        ws.cell(row=1, column=c).border = THIN_BORDER
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 50
    ws.column_dimensions["D"].width = 30
    ws.column_dimensions["E"].width = 10
    ws.column_dimensions["F"].width = 30
    row = 2
    for entry in audit_trail or []:
        if isinstance(entry, dict):
            ts = entry.get("timestamp_utc", "")
            event = entry.get("event_type", "")
            reasoning = entry.get("reasoning", "")
            citations = entry.get("citations", "")
            outcome = entry.get("outcome", "")
            warning = entry.get("warning_message") or ""
        else:
            ts = getattr(entry, "timestamp_utc", "")
            event = getattr(entry, "event_type", "")
            reasoning = getattr(entry, "reasoning", "")
            citations = getattr(entry, "citations", "")
            outcome = getattr(entry, "outcome", "")
            warning = getattr(entry, "warning_message") or ""
        ws.cell(row=row, column=1, value=ts)
        ws.cell(row=row, column=2, value=event)
        ws.cell(row=row, column=3, value=reasoning)
        ws.cell(row=row, column=3).alignment = Alignment(wrap_text=True)
        ws.cell(row=row, column=4, value=citations)
        ws.cell(row=row, column=5, value=outcome)
        ws.cell(row=row, column=6, value=warning)
        row += 1
    if row == 2:
        ws.cell(row=2, column=1, value="No audit trail entries in this report.")


def build_excel(payload: dict[str, Any]) -> bytes:
    """
    Build a Big-4 style Excel workbook from a report payload (dict or ReportPayload-like).
    Uses formulas for totals (SUM) and net income (revenue - expenses) where possible.
    Returns Excel file bytes.
    """
    wb = Workbook()
    cover = _get_cover(payload)
    _write_cover_sheet(wb.active, cover)
    financials = _get_financials(payload)
    ws_fs = wb.create_sheet("Financial Statements", 1)
    _write_financial_sheet(ws_fs, financials)
    strategic = payload.get("strategic_analysis") or ""
    ws_sa = wb.create_sheet("Strategic Analysis", 2)
    _write_strategic_sheet(ws_sa, strategic)
    audit = payload.get("audit_trail") or []
    ws_at = wb.create_sheet("Audit Trail", 3)
    _write_audit_sheet(ws_at, audit)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()
